import base64
import json
import os
import re
from pathlib import Path
from typing import Callable

import httpx
from openpyxl import Workbook, load_workbook

from skills.base import BaseSkill
from services.image_backend import get_backend, read_image
from services.llm import sanitize_image_prompt_for_safety

# ── Claude / LLM ─────────────────────────────────────────────────────────────
API_KEY = os.getenv("ANTHROPIC_API_KEY", "")
API_BASE_URL = os.getenv("API_BASE_URL", "https://api.anthropic.com/v1").rstrip("/")
CHAT_URL = f"{API_BASE_URL}/chat/completions"
CHAT_MODEL = "claude-sonnet-4-6"

OUTPUT_BASE = Path("outputs/sketch2portrait")
OUTPUT_BASE.mkdir(parents=True, exist_ok=True)

DEFAULT_MAIN_PROMPT = (
    "参考图包含角色正面与背面及道具信息。\n"
    "纯白背景，单角色单视角2D立绘，保持外形、比例与道具一致。\n"
    "卡通动漫描边（外轮廓略粗，内部线细），干净上色。\n"
    "全身完整入镜，四周留白。"
)


class Sketch2PortraitSkill(BaseSkill):

    @classmethod
    def skill_id(cls) -> str:
        return "sketch2portrait"

    @classmethod
    def display_name(cls) -> str:
        return "原画转2D"

    @classmethod
    def description(cls) -> str:
        return "上传参考图，AI 分析角色并批量生成高质量 2D 动作插画"

    @classmethod
    def input_schema(cls) -> dict:
        return {
            "type": "object",
            "properties": {
                "image_path": {"type": "string", "title": "参考图路径"},
                "task_name": {"type": "string", "title": "任务名字"},
                "actions_count": {
                    "type": "integer", "title": "动作组数",
                    "default": 4, "minimum": 1, "maximum": 10,
                },
                "main_prompt": {"type": "string", "title": "主提示词"},
                "image_backend": {
                    "type": "string", "title": "生图模型",
                    "enum": ["seedream", "gemini", "gpt-image-2"], "default": "seedream",
                },
            },
            "required": ["image_path", "task_name", "actions_count"],
        }

    def execute(
        self,
        params: dict,
        on_progress: Callable[[int, str], None],
        runtime_context: dict | None = None,
    ) -> dict:
        image_path = params.get("image_path", "").strip()
        task_name = params.get("task_name", "").strip()
        actions_count = int(params.get("actions_count", 4))
        main_prompt = params.get("main_prompt", "").strip() or DEFAULT_MAIN_PROMPT
        image_backend = params.get("image_backend", "seedream")
        safety_rewrite_prompt = bool(params.get("safety_rewrite_prompt", False))

        if not image_path:
            raise ValueError("缺少 image_path")
        if not task_name:
            raise ValueError("缺少 task_name")
        if actions_count < 1:
            raise ValueError("actions_count 必须 >= 1")

        on_progress(10, "正在初始化任务工作区...")
        ws_paths = self._init_workspace(task_name, image_path, actions_count, main_prompt)

        on_progress(25, "AI 正在分析参考图并生成动作模板...")
        templates = self._generate_action_templates(image_path, actions_count, main_prompt)

        on_progress(50, "正在写入动作草稿...")
        self._write_draft(ws_paths, task_name, image_path, main_prompt, templates)

        on_progress(60, "自动审核通过，写入 review...")
        self._write_review(ws_paths, task_name, templates)

        on_progress(70, "正在写回批量表格...")
        self._write_batch_xlsx(ws_paths, templates)

        on_progress(75, f"提交 {image_backend} 批量生图...")
        saved = self._run_batch(
            ws_paths, image_path, main_prompt, templates, image_backend, safety_rewrite_prompt, on_progress
        )

        return {
            "task_root": str(ws_paths["task_root"]),
            "output_dir": str(ws_paths["result_dir"]),
            "output_files": saved,
            "draft_path": str(ws_paths["draft_path"]),
            "batch_xlsx": str(ws_paths["batch_xlsx"]),
        }

    # ── workspace ─────────────────────────────────────────────────────────────
    def _init_workspace(self, task_name, image_path, actions_count, main_prompt) -> dict:
        safe = re.sub(r'[\\/:*?"<>|]+', "_", task_name).strip("._") or "task"
        task_root = OUTPUT_BASE / safe
        draft_dir = task_root / "01_输入与草稿"
        table_dir = task_root / "02_写回表格"
        result_dir = task_root / "03_批量结果"
        for d in [draft_dir, table_dir, result_dir]:
            d.mkdir(parents=True, exist_ok=True)

        source_xlsx = table_dir / "01_source.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "batch"
        for col, h in enumerate(["image", "main_prompt", "action_template"], 1):
            ws.cell(1, col, h)
        img_abs = str(Path(image_path).resolve())
        for i in range(actions_count):
            ws.cell(2 + i, 1, img_abs)
            ws.cell(2 + i, 2, main_prompt)
        wb.save(source_xlsx)

        return {
            "task_root": task_root, "draft_dir": draft_dir,
            "table_dir": table_dir, "result_dir": result_dir,
            "source_xlsx": source_xlsx,
            "batch_xlsx": table_dir / "04_batch.xlsx",
            "draft_path": draft_dir / "01_draft.md",
            "review_path": draft_dir / "02_review.md",
        }

    # ── xlsx helpers ──────────────────────────────────────────────────────────
    def _write_draft(self, ws_paths, task_name, image_path, main_prompt, templates):
        items = "\n".join(
            f"### item\naction_id: A{i+1:03d}\naction_name: {t['action_name']}\n"
            f"action_template: |\n  {t['action_template']}\n"
            for i, t in enumerate(templates)
        )
        ws_paths["draft_path"].write_text(
            f"# Action Template Draft\n\ntask_name: {task_name}\n"
            f"reference_image: {image_path}\nsource_xlsx: {ws_paths['source_xlsx']}\n"
            f"main_prompt: |\n  {main_prompt.replace(chr(10), chr(10)+'  ')}\n\n---\n\n## items\n\n{items}",
            encoding="utf-8",
        )

    def _write_review(self, ws_paths, task_name, templates):
        items = "\n".join(
            f"### item\naction_id: A{i+1:03d}\naction_name: {t['action_name']}\n"
            f"action_template: |\n  {t['action_template']}\n"
            for i, t in enumerate(templates)
        )
        ws_paths["review_path"].write_text(
            f"# 动作模板审核\n\ntask_name: {task_name}\n"
            f"review_scope: whole_group\ngroup_review_status: 赞同\n"
            f"group_review_note: 自动审核通过\n\n---\n\n## 条目\n\n{items}",
            encoding="utf-8",
        )

    def _write_batch_xlsx(self, ws_paths, templates):
        import shutil
        shutil.copy(ws_paths["source_xlsx"], ws_paths["batch_xlsx"])
        wb = load_workbook(ws_paths["batch_xlsx"])
        ws = wb["batch"]
        headers = [(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
        col = headers.index("action_template") + 1 if "action_template" in headers else ws.max_column + 1
        if "action_template" not in headers:
            ws.cell(1, col, "action_template")
        for i, t in enumerate(templates):
            ws.cell(2 + i, col, t["action_template"])
        wb.save(ws_paths["batch_xlsx"])

    # ── image generation ──────────────────────────────────────────────────────
    def _run_batch(self, ws_paths, image_path, main_prompt, templates,
                   image_backend: str, safety_rewrite_prompt: bool, on_progress: Callable) -> list[str]:
        backend = get_backend(image_backend)
        ref_bytes, ref_mime = read_image(image_path)
        result_dir: Path = ws_paths["result_dir"]
        task_name = result_dir.parent.name
        saved, action_log = [], []
        total = len(templates)

        for idx, t in enumerate(templates):
            pct = 75 + int(idx / total * 20)
            on_progress(pct, f"{image_backend} 生图 {idx+1}/{total}：{t['action_name']}...")

            full_prompt = f"{main_prompt}\n{t['action_template']}"
            if image_backend == "gpt-image-2" and safety_rewrite_prompt:
                full_prompt = sanitize_image_prompt_for_safety(
                    user_prompt=full_prompt,
                    has_reference=True,
                    target_model="gpt-image-2",
                    refusal_message="manual safety retry requested by user",
                )["sanitized_prompt"].strip()
            img_bytes = backend.generate(
                prompt=full_prompt,
                ref_images=[ref_bytes],
                ref_mimetypes=[ref_mime],
                size="2048x2048",
                aspect_ratio="1:1",
            )
            out_path = result_dir / f"A{idx+1:03d}_{t['action_name']}.png"
            out_path.write_bytes(img_bytes)
            saved.append(str(out_path))
            action_log.append({"index": idx + 1, "file": out_path.name,
                                "action_name": t["action_name"],
                                "action_template": t["action_template"],
                                "full_prompt": full_prompt})

        from datetime import datetime, timezone
        (result_dir / "run.json").write_text(json.dumps({
            "skill": "sketch2portrait", "task_name": task_name,
            "image_backend": image_backend, "actions_count": total,
            "manual_safety_retry": safety_rewrite_prompt,
            "main_prompt": main_prompt, "reference_image": image_path,
            "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "actions": action_log,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        return saved

    # ── Claude action template generation ────────────────────────────────────
    def _generate_action_templates(self, image_path, actions_count, main_prompt) -> list:
        if not API_KEY:
            raise RuntimeError("未配置 ANTHROPIC_API_KEY，无法调用 AI 分析参考图")

        ref_bytes, ref_mime = read_image(image_path)
        img_b64 = base64.b64encode(ref_bytes).decode()

        user_prompt = (
            f"你是一个角色动作设计专家。请分析参考图中的角色，生成 {actions_count} 个 action_template。\n\n"
            "规则：\n"
            "1. 严格保持角色一致性——只使用参考图中可见的特征（外貌/道具/着装），不新增未出现的特征\n"
            "2. 若角色是动物融合设计，禁止按常识补全参考图未出现的器官\n"
            "3. 每条 action_template 必须包含：镜头机位+画面畸变+身体夸张形变但不离谱+招式动作+无特效+卡通游戏高级感\n"
            "4. 白底单角色场景，聚焦动作可读性\n"
            "5. 生成后自检：若出现参考图中不存在的特征词，必须修改\n"
            "6. 严格规则：所有字符串值内不得出现英文双引号或换行符，用顿号替代列举\n\n"
            f"主提示词参考：{main_prompt}\n\n"
            "请以 JSON 数组形式返回（不要 markdown 包裹），格式：\n"
            '[{"action_name": "动作名2到6字", "action_template": "一段话中文60到120字单行"}, ...]\n\n'
            f"生成 {actions_count} 个不同招式动作。"
        )

        resp = httpx.post(
            CHAT_URL,
            headers={"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"},
            json={
                "model": CHAT_MODEL,
                "max_tokens": 2048,
                "thinking": {"type": "disabled"},
                "messages": [{"role": "user", "content": [
                    {"type": "image_url", "image_url": {"url": f"data:{ref_mime};base64,{img_b64}"}},
                    {"type": "text", "text": user_prompt},
                ]}],
            },
            timeout=120,
        )
        resp.raise_for_status()

        raw = resp.json()["choices"][0]["message"]["content"].strip()
        # Strip code fences
        if "```json" in raw:
            raw = raw.split("```json", 1)[1].split("```", 1)[0].strip()
        elif "```" in raw:
            raw = max((p.strip() for p in raw.split("```")), key=len)
        # Extract JSON array if preamble text present
        if not raw.startswith("["):
            m = re.search(r'\[[\s\S]*\]', raw)
            if m:
                raw = m.group(0)
        try:
            templates = json.loads(raw)
        except json.JSONDecodeError:
            # Fallback: extract individual objects
            objects = re.findall(r'\{[\s\S]*?\}(?=\s*[,\]])', raw)
            templates = []
            for obj in objects:
                try:
                    templates.append(json.loads(obj))
                except json.JSONDecodeError:
                    name = re.search(r'"action_name"\s*:\s*"([^"]+)"', obj)
                    tpl = re.search(r'"action_template"\s*:\s*"([\s\S]+?)"(?=\s*[,}])', obj)
                    if name and tpl:
                        templates.append({"action_name": name.group(1), "action_template": tpl.group(1)})
        if not isinstance(templates, list) or not templates:
            raise ValueError("AI 返回的 action_template 格式不正确")
        return templates
